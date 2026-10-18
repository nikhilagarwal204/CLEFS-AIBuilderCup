"""
Demo seed, money part: Profit Guard statements, payouts, cards, bank, folio
and every planted leakage case (02 spec 9 "Statements" and "Planted cases").

Everything enters through the real pipeline: each document is rendered as
the file an owner would download (Booking CSV in three header variants, an
Agoda XLSX, an Expedia Hotel Collect PDF made with fpdf2, a credit note,
the Booking payouts export, a bank CSV, VCC reports and a PMS folio export),
stored with clefs.files, staged by clefs.core.ingest and committed through
the Profit Guard commit handlers. Detection then runs the registered job and
task bodies (clefs.jobs registry: pg.detect_on_commit, pg.auto_verify,
tax.td_ledger_rollup, tax.td_checks) inside the seed transaction, each with
the clock frozen at the moment a worker would have run it:

    Jul 3   June statement (v1 CSV)        -> detect (L1b: fees not configured yet)
    Aug 3   July statement (v2 UTF-16 TSV)
    Sep 3   August statement (v3 ; CSV)    -> detect (L3 before the credit note)
    Sep 5-6 Agoda XLSX, Expedia HC PDF
    Sep 12  credit note                    -> auto-verify (2 L3 findings RECOVERED)
    A-2     payouts; A-1 bank CSV and PMS folio export; A 08:15 VCC reports
    the 1st of each month up to A: TD ledger close (tax.td_ledger_rollup),
            remitted on the 10th (August AED 350 short)
    A 21:00 full detector run (pg.detect_on_commit, kind manual) + tax.td_checks

(A = the anchor. Statement months are fixed at Jun-Aug 2026 because the
programs and fee configuration they test are; every other plant moves with
the anchor. Anchors before 23 Sep 2026 skip this part.)

Why dedicated plant stays (pms ids RC9xxxxx) rather than bending core
stays: every planted case must exist at any --scale, with exact amounts,
and core_stays is deliberately clean (no early departures, no OTA
no-shows). Plant stays get the same night-audit folio as core stays, so
they raise nothing but the case they plant.

Why a replay AI backend for the PDF: ingest reads PDFs with pypdf (real)
and then asks clefs.ai to map the text to rows. The seed never calls the
network, so for that one call it installs a backend that answers with the
rows the PDF was rendered from (the same shape a model returns), then puts
the previous backend back.

Why the Booking payouts JSON is imported as CSV: core/ingest reads tables
(CSV/TSV/XLSX/XLS) and PDFs only; a JSON file is refused (415). The JSON
(the Payouts API shape) is kept as the raw artifact and flattened to the
extranet's CSV export, which is what gets imported.

Data this part adjusts besides its own rows (all hotel-scoped except one):
the two GENIUS programs core_catalog seeds get rate plans (GENIUS_L1 /
GENIUS_L2+L3) so one level never stacks on the other; L10's expected bank
fee (pfg_detector_tuning, AED 25); the owner's override applying the
announced TD exemption for one night (hotel_tax_overrides); and, only when
fx_rates has no EUR/AED row at all, one MANUAL reference rate (L11 needs it).

Plants go to ctx.plants: {detector, key, klass, currency, expected_minor,
expected, fingerprints, subjects, ids, decoy_ids, decoy_fingerprints, note}.
klass is the finding class, or DECOY for an entry that only lists decoys,
or STRUCTURAL for Agoda's by-design records (never claims). A fingerprint
is Profit Guard's own sha1(detector|ota|res_norm|doc); the tax checks key
by database id, so their plants carry subjects [{pms_res_id, month}] and
fingerprints_of() resolves them. No database id is written, so the same
seed and anchor give the same manifest bytes. Expected amounts are the
detectors' own arithmetic on the planted data (exact at the canonical run).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Iterator
from zoneinfo import ZoneInfo

from clefs.core import clock
from clefs.core import stays as core_stays
from clefs.core import tax as core_tax

PART_ORDER = 50
UTC = ZoneInfo("UTC")
CCY = "AED"

# ─── Calendar of the money story ───

# The statements cover the three months the spec names. They are absolute
# because the programs and the Booking fee configuration they test are
# absolute too (core_catalog: Preferred 1 Jun-31 Jul 2026, MF/TD configured
# as charges from 1 Aug 2026).
JUN, JUL, AUG, SEP = date(2026, 6, 1), date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1)
STATEMENT_MONTHS = (JUN, JUL, AUG)
# The part needs the August statement and the credit note to have arrived,
# and the exempt night (anchor - 7) to fall after the T2 chain (to Sep 16).
EARLIEST_ANCHOR = date(2026, 9, 23)
# The night the owner applied the announced TD exemption to (TX1): after
# every September plant stay, before the L14 no-shows.
EXEMPT_NIGHT_OFFSET = 7          # anchor - 7 days
TD_RULE = "dxb_2026_exempt_td"
BANK_FEE_MINOR = 2500            # AED 25 per incoming transfer (L10's expected fee)
TD_SHORT_MINOR = 35000           # August TD remitted AED 350 short (T7)
PLANT_ID_BASE = 900000           # plant stays are RC900001...
FILES_ENV = "CLEFS_DEMO_FILES_DIR"


def month_end(d: date) -> date:
    nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return nxt - timedelta(days=1)


def mkey(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def rnd(x: Any) -> int:
    return int(Decimal(str(x)).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def major(minor: int) -> str:
    """Minor units as a Decimal string with two places ("1240.00")."""
    return str((Decimal(int(minor)) / 100).quantize(Decimal("0.01")))


def fingerprint(detector: str, ota: str | None, res_norm: str | None, doc: str | None) -> str:
    """Profit Guard's finding fingerprint (detectors/base.py), recomputed
    here so the manifest can name the exact finding a plant must produce."""
    raw = "|".join([detector, ota or "", res_norm or "", doc or ""])
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


# ─── State of one run ───


@dataclass
class Money:
    """What this part made, for the later steps and the manifest."""

    ctx: Any
    api: Any                                             # clefs.modules.profit_guard.api
    rng: Any = None                                      # random.Random for ids and amounts
    stays: dict[str, dict] = field(default_factory=dict)  # plant key -> stay
    files: dict[str, bytes] = field(default_factory=dict)
    batches: dict[str, dict] = field(default_factory=dict)
    lines: dict[str, list[dict]] = field(default_factory=dict)  # document -> printed lines
    next_id: int = PLANT_ID_BASE
    ota_ids: set[str] = field(default_factory=set)
    counts: dict[str, int] = field(default_factory=dict)

    @property
    def cur(self) -> Any:
        return self.ctx.cur

    @property
    def slug(self) -> str:
        return self.ctx.slug

    def bump(self, key: str, n: int = 1) -> None:
        self.counts[key] = self.counts.get(key, 0) + n


@contextmanager
def at(ctx: Any, d: date, hour: int = 10, minute: int = 0) -> Iterator[datetime]:
    """Run a step as a worker would have run it then: the clock is frozen at
    that hotel-local instant and set back to the anchor evening after."""
    instant = ctx.at(d, hour, minute)
    clock.freeze(instant.astimezone(UTC))
    try:
        yield instant
    finally:
        clock.freeze(ctx.now.astimezone(UTC))


def plant(m: Money, detector: str, key: str, *, klass: str, expected_minor: int, fingerprints: list[str],
          ids: dict, decoy_ids: dict | None = None, decoy_fingerprints: list[str] | None = None,
          note: str | None = None, currency: str = CCY, subjects: list[dict] | None = None) -> None:
    """One manifest entry. `fingerprints` are the findings the case must
    produce (one per booking, card, payout or month); the decoys must
    produce none of theirs. The tax checks key a finding by the database
    id of the reservation (res:<id>), which changes on every re-seed, so
    those plants name `subjects` [{pms_res_id, month}] instead and
    fingerprints_of() resolves them: the manifest stays byte-stable."""
    m.ctx.plants.append({
        "detector": detector, "key": key, "klass": klass, "currency": currency,
        "expected_minor": int(expected_minor), "expected": major(expected_minor),
        "fingerprints": sorted(fingerprints), "subjects": subjects or [], "ids": ids, "decoy_ids": decoy_ids or {},
        "decoy_fingerprints": sorted(decoy_fingerprints or []), "note": note,
    })


def fingerprints_of(cur: Any, slug: str, p: dict) -> list[str]:
    """Every fingerprint a manifest entry expects, its subjects resolved
    against this database (tests and demo-check call this)."""
    out = list(p.get("fingerprints") or [])
    for sub in p.get("subjects") or []:
        cur.execute("SELECT id FROM reservations WHERE hotel_slug = %s AND source_system = 'PMS' AND pms_res_id = %s",
                    (slug, sub["pms_res_id"]))
        row = cur.fetchone()
        out.append(fingerprint(p["detector"], None, f"res:{row['id'] if row else 'missing'}", sub["month"]))
    return out


# ─── Plant stays (through core.stays) ───

# Guests of the plant stays: (name, country, language). Plain names, as a
# PMS holds them; the hotel itself carries the "(DEMO)" marker.
GUESTS = (
    ("Omar Haddad", "JO", "ar"), ("Priya Menon", "IN", "en"), ("Lukas Brandt", "DE", "en"),
    ("Aisha Rahman", "AE", "ar"), ("Dmitri Volkov", "RU", "ru"), ("Chen Wei", "CN", "zh"),
    ("Fatima Al Suwaidi", "AE", "ar"), ("James Carter", "GB", "en"), ("Anjali Sharma", "IN", "hi"),
    ("Youssef Benali", "MA", "ar"), ("Elena Petrova", "RU", "ru"), ("Hassan Qureshi", "PK", "en"),
    ("Sofia Rossi", "IT", "en"), ("Karim Nasser", "LB", "ar"), ("Mei Lin", "SG", "zh"),
    ("Rahul Verma", "IN", "hi"), ("Sara Khalil", "EG", "ar"), ("Tom Becker", "NL", "en"),
)
ROOM_NAMES = {"STD_KING": "Standard King Room", "STD_TWIN": "Standard Twin Room",
              "FAM_2BR": "Two-Bedroom Family Suite"}
OTA_ID_RANGE = {"BOOKING": (4_000_000_000, 4_999_999_999), "AGODA": (800_000_000, 999_999_999),
                "EXPEDIA": (7_100_000_000, 7_399_999_999)}
FOLIO_CODES = (("rate_minor", "ROOM"), ("sc_minor", "SC"), ("mf_minor", "MF"), ("tf_minor", "TF"),
               ("vat_minor", "VAT"), ("td_minor", "TD"))


def _ota_id(m: Money, channel: str) -> str | None:
    """A fresh OTA reservation id, never one a core stay already holds (the
    matcher links by it)."""
    if channel not in OTA_ID_RANGE:
        return None
    lo, hi = OTA_ID_RANGE[channel]
    while True:
        v = str(m.rng.randint(lo, hi))
        if v not in m.ota_ids:
            m.ota_ids.add(v)
            return v


def make_stay(m: Money, key: str, *, channel: str, check_in: date, nights: int, rate: int,
              status: str = "CHECKED_OUT", collect_model: str | None = "HOTEL_COLLECT", rooms: int = 1,
              room_type: str = "STD_KING", actual_nights: int | None = None, rate_plan: str = "BAR",
              no_show_fee: int | None = None, cancel_fee: int | None = None, guest: int | None = None,
              audit: bool = True, rates: list[int] | None = None) -> dict:
    """One plant stay through upsert_reservation (nights, fees, chains and
    stay events from the single write path). `rate` is AED per room-night;
    `actual_nights` < nights makes an early departure; `audit` posts the
    night audit's folio for the stayed nights (as core_stays does)."""
    ctx = m.ctx
    m.next_id += 1
    name, country, lang = GUESTS[(guest if guest is not None else m.next_id) % len(GUESTS)]
    co = check_in + timedelta(days=nights)
    per_night = rates or [rate] * nights
    data: dict[str, Any] = {
        "pms_res_id": f"RC{m.next_id}", "channel_code": channel, "ota_res_id": _ota_id(m, channel),
        "guest_name": name, "guest_country": country, "guest_language": lang,
        "booked_at": ctx.at(check_in, 11) - timedelta(days=21 + m.next_id % 17), "check_in": check_in,
        "check_out": co, "rooms_count": rooms, "room_type_code": room_type, "rate_plan_code": rate_plan,
        "status": status, "collect_model": collect_model, "currency": CCY, "adults": 2,
        "payment_method": "VCC" if collect_model == "OTA_COLLECT" else None,
    }
    if status in ("CHECKED_OUT", "IN_HOUSE"):
        data["checked_in_at"] = ctx.at(check_in, 15, 20)
    if status == "CHECKED_OUT":
        left = check_in + timedelta(days=actual_nights if actual_nights is not None else nights)
        if left < co:
            data["actual_check_out"] = left
        data["checked_out_at"] = ctx.at(left, 10, 5)
    if status == "CANCELLED":
        data["cancelled_at"] = ctx.at(check_in - timedelta(days=2), 18)
        data["cancel_fee_minor"] = cancel_fee or 0
    if status == "NO_SHOW":
        data["no_show_fee_minor"] = no_show_fee or 0
    night_rows = [{"stay_date": check_in + timedelta(days=i), "rate_minor": r * 100, "room_type_code": room_type}
                  for i, r in enumerate(per_night)]
    res = core_stays.upsert_reservation(m.cur, m.slug, data, source_system="PMS", nights=night_rows)
    stay = data | {"id": int(res["id"]), "key": key, "nights": nights, "rates": per_night,
                   "norm": core_stays.normalize_ota_res_id(channel, data["ota_res_id"])
                   if data["ota_res_id"] else None}
    m.stays[key] = stay
    if audit:
        post_folio(m, folio_rows(m, [stay["id"]]))
    return stay


def folio_rows(m: Money, reservation_ids: list[int]) -> list[dict]:
    """The night audit's postings for the stayed nights of these stays: the
    same codes, ids and amounts core_stays posts for core stays."""
    m.cur.execute(
        """
        SELECT r.pms_res_id, n.stay_date, n.room_number, n.rooms, n.rate_minor, n.sc_minor, n.mf_minor,
               n.tf_minor, n.vat_minor, n.td_minor
          FROM reservation_nights n JOIN reservations r ON r.id = n.reservation_id
         WHERE n.hotel_slug = %s AND n.stayed AND r.id = ANY(%s)
         ORDER BY r.pms_res_id, n.stay_date, n.room_seq
        """,
        (m.slug, reservation_ids),
    )
    rows: list[dict] = []
    for n in m.cur.fetchall():
        tag = f"{n['pms_res_id']}-{n['stay_date']:%Y%m%d}"
        for col, code in FOLIO_CODES:
            amount = int(n[col] or 0) * (int(n["rooms"] or 1) if col == "rate_minor" else 1)
            if amount:
                rows.append({"pms_txn_id": f"{tag}-{code}", "posted_date": n["stay_date"],
                             "stay_date": n["stay_date"], "code": code, "amount_minor": amount,
                             "pms_res_id": n["pms_res_id"], "room_number": n["room_number"],
                             "description": code.title(), "qty": 1})
    return rows


def post_folio(m: Money, rows: list[dict]) -> None:
    if rows:
        core_stays.upsert_folio_charges(m.cur, m.slug, rows, currency=CCY)
        m.bump("folio_charges", len(rows))


# ─── Commission views ───
# A statement line is printed from a "view": what the OTA believes the stay
# cost (gross, commissionable base, percent, commission). Clean lines use
# the PMS's own breakdown, so they agree with the detectors to the fil; a
# plant uses a different view (the booked nights, a pre-modification rate).


def _view(m: Money, ota: str, breakdown: Any, on: date) -> dict:
    exp = m.api.expected_commission(m.cur, m.slug, ota_code=ota, breakdown=breakdown, on=on)
    if exp is None:
        raise RuntimeError(f"money seed: no commission terms for {ota} on {on}")
    return {"gross": int(breakdown.total), "base": int(exp["base_minor"]), "commission": int(exp["commission_minor"]),
            "pct": Decimal(str(exp["pct"])), "mf": int(breakdown.mf), "td": int(breakdown.td),
            "vat": int(breakdown.vat), "room": int(breakdown.room)}


def pms_view(m: Money, ota: str, reservation_id: int, on: date) -> dict:
    return _view(m, ota, core_stays.reservation_breakdown(m.cur, reservation_id), on)


def quoted_view(m: Money, ota: str, check_in: date, rates: list[int], room_type: str = "STD_KING",
                rooms: int = 1) -> dict:
    """The view of a stay at these nightly rates (AED), composed through
    core.tax the way the PMS composes its own nights."""
    lines = [core_tax.StayLine(stay_date=check_in + timedelta(days=i), room_type_code=room_type, rooms=rooms,
                               rate_minor=r * 100) for i, r in enumerate(rates)]
    return _view(m, ota, core_tax.compose_stay(m.cur, m.slug, lines=lines), check_in)


def pct_of(amount: int, pct: Decimal) -> int:
    return rnd(Decimal(amount) * pct / 100)


def make_line(r: dict, v: dict, *, status: str = "OK", **over: Any) -> dict:
    """One printed statement line from a reservation row and a view.
    Keys: ota_res_id, guest, check_in, check_out, nights, rooms, status,
    gross, base, pct, commission, country, rate_plan, room_type, programs,
    booked_on, collect_model, remarks, reservation_id."""
    ln = {
        "ota_res_id": r["ota_res_id"], "guest": r["guest_name"], "check_in": r["check_in"],
        "check_out": r["check_out"], "nights": (r["check_out"] - r["check_in"]).days,
        "rooms": int(r.get("rooms_count") or 1), "status": status, "gross": v["gross"], "base": v["base"],
        "pct": v["pct"], "commission": v["commission"], "country": r.get("guest_country") or "",
        "rate_plan": r.get("rate_plan_code") or "BAR", "room_type": ROOM_NAMES.get(r.get("room_type_code") or "",
                                                                                    "Standard Room"),
        "programs": "", "booked_on": (r.get("booked_at") or datetime.combine(r["check_in"], time(12))).date(),
        "collect_model": r.get("collect_model") or "HOTEL_COLLECT", "remarks": "",
        "reservation_id": r.get("id"), "mf": v.get("mf", 0), "td": v.get("td", 0), "vat": v.get("vat", 0),
    }
    ln.update(over)
    return ln


_CORE_SQL = """
    SELECT id, pms_res_id, ota_res_id, guest_name, guest_country, check_in, check_out, rooms_count, status,
           rate_plan_code, room_type_code, booked_at, collect_model
      FROM reservations
     WHERE hotel_slug = %s AND channel_code = %s AND source_system = 'PMS' AND status = 'CHECKED_OUT'
       AND actual_check_out IS NULL AND check_out BETWEEN %s AND %s AND pms_res_id NOT LIKE 'RC9%%'
       AND (%s::text IS NULL OR collect_model = %s)
     ORDER BY check_out, pms_res_id
"""


def core_lines(m: Money, ota: str, month: date, *, collect_model: str | None = None) -> list[dict]:
    """Clean lines: every core stay of the OTA that checked out in the
    month (the statement period), priced from the PMS's own breakdown."""
    m.cur.execute(_CORE_SQL, (m.slug, ota, month, month_end(month), collect_model, collect_model))
    rows = [dict(r) for r in m.cur.fetchall()]
    return [make_line(r, pms_view(m, ota, r["id"], r["check_in"])) for r in rows]


def stay_row(stay: dict) -> dict:
    """A plant stay in the shape make_line reads."""
    return {"id": stay["id"], "ota_res_id": stay["ota_res_id"], "guest_name": stay["guest_name"],
            "guest_country": stay["guest_country"], "check_in": stay["check_in"], "check_out": stay["check_out"],
            "rooms_count": stay["rooms_count"], "rate_plan_code": stay["rate_plan_code"],
            "room_type_code": stay["room_type_code"], "booked_at": stay["booked_at"],
            "collect_model": stay["collect_model"]}


# ─── Rendering: Booking.com reservation statements ───
# Three layouts of the same export, as Booking has shipped over the years
# (UNVERIFIED headers, core/ingest aliases/booking_res_statement.yaml): the
# alias map must read all three without a saved template.


def dot(minor: int) -> str:
    return f"{Decimal(int(minor)) / 100:.2f}"


def comma(minor: int) -> str:
    """European money: "1.234,56"."""
    return f"{Decimal(int(minor)) / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct_txt(p: Decimal, *, dec: str = ".") -> str:
    return f"{p:.2f}".replace(".", dec)


def _csv(rows: list[list[Any]], *, delimiter: str = ",") -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    for r in rows:
        w.writerow(["" if c is None else c for c in r])
    return buf.getvalue()


BOOKING_STATUS = {"OK": "ok", "NO_SHOW": "no_show", "CANCELLED": "cancelled_by_guest"}
BOOKING_V1 = ("Book number", "Guest name", "Check-in", "Check-out", "Booked on", "Status", "Rooms", "Room nights",
              "Price", "Commission %", "Commission amount", "Commissionable amount", "Currency", "Booker country",
              "Rate name", "Unit type", "Programs", "Remarks", "Doc no")
BOOKING_V2 = ("Reservation number", "Guest names", "Arrival", "Departure", "Booking date", "Reservation status",
              "Number of rooms", "Duration (nights)", "Final amount", "Commission percentage", "Commission",
              "Commissionable revenue", "Currency code", "Guest country", "Rate plan", "Room name", "Programs",
              "Notes", "Doc no")
BOOKING_V3 = ("Booking number", "Booker name", "Check in", "Check out", "Reservation date", "Status", "Rooms",
              "Nights", "Total price", "Commission rate", "Commission amount payable", "Commissionable price",
              "Currency", "Country", "Policy", "Room type", "Programs", "Remarks", "Doc no")


def _booking_cells(ln: dict, *, money: Any, day: Any, dec: str, doc_no: str) -> list[Any]:
    return [ln["ota_res_id"], ln["guest"], day(ln["check_in"]), day(ln["check_out"]), day(ln["booked_on"]),
            BOOKING_STATUS[ln["status"]], ln["rooms"], ln["nights"] * ln["rooms"], money(ln["gross"]),
            pct_txt(ln["pct"], dec=dec), money(ln["commission"]), money(ln["base"]), CCY, ln["country"],
            ln["rate_plan"], ln["room_type"], ln["programs"], ln["remarks"], doc_no]


def _total_row(width: int, at_col: int, total: str) -> list[Any]:
    row: list[Any] = [""] * width
    row[0], row[at_col] = "Total", total
    return row


def render_booking(lines: list[dict], *, variant: int, doc_no: str) -> tuple[bytes, str, str]:
    """(bytes, filename, content type) of one monthly reservation statement.
    v1: UTF-8 CSV, ISO dates. v2: UTF-16 TSV (the Excel "Unicode text"
    save), day-first dates. v3: semicolon CSV with decimal commas and dotted
    dates (a European-locale Excel)."""
    total = sum(ln["commission"] for ln in lines)
    if variant == 1:
        rows = [list(BOOKING_V1)] + [_booking_cells(ln, money=dot, day=date.isoformat, dec=".", doc_no=doc_no)
                                     for ln in lines]
        rows.append(_total_row(len(BOOKING_V1), 10, dot(total)))
        return _csv(rows).encode("utf-8"), f"booking-statement-{doc_no}.csv", "text/csv"
    if variant == 2:
        rows = [list(BOOKING_V2)] + [_booking_cells(ln, money=dot, day=lambda d: d.strftime("%d/%m/%Y"), dec=".",
                                                    doc_no=doc_no) for ln in lines]
        rows.append(_total_row(len(BOOKING_V2), 10, dot(total)))
        return (_csv(rows, delimiter="\t").encode("utf-16"), f"booking-statement-{doc_no}.txt",
                "text/tab-separated-values")
    rows = [list(BOOKING_V3)] + [_booking_cells(ln, money=comma, day=lambda d: d.strftime("%d.%m.%Y"), dec=",",
                                                doc_no=doc_no) for ln in lines]
    rows.append(_total_row(len(BOOKING_V3), 10, comma(total)))
    return _csv(rows, delimiter=";").encode("utf-8-sig"), f"booking-statement-{doc_no}.csv", "text/csv"


# ─── Rendering: Agoda XLSX, Expedia Hotel Collect PDF ───

AGODA_HEADERS = ("Booking ID", "Customer Name", "Check In Date", "Check Out Date", "Booking Date", "Status",
                 "Room Nights", "No of Rooms", "Room Type", "Rate Plan", "Payment Model", "Reference Sell Rate",
                 "Commissionable Amount", "Commission Rate", "Commission", "Currency", "Doc no")
AGODA_STATUS = {"OK": "Departed", "NO_SHOW": "No Show", "CANCELLED": "Cancelled"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIXED_TS = datetime(2026, 9, 5, 8, 0)  # file metadata only: identical bytes on every run


def render_agoda(lines: list[dict], *, doc_no: str, month: date) -> tuple[bytes, str, str]:
    """Agoda's finance report as the XLSX the YCS portal exports: a title,
    a blank row, then the table and a Total row (real dates and numbers in
    the cells, as Excel keeps them)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append([f"Agoda finance report {mkey(month)}"])
    ws.append([])
    ws.append(list(AGODA_HEADERS))
    for ln in lines:
        ws.append([ln["ota_res_id"], ln["guest"], ln["check_in"], ln["check_out"], ln["booked_on"],
                   AGODA_STATUS[ln["status"]], ln["nights"] * ln["rooms"], ln["rooms"], ln["room_type"],
                   ln["rate_plan"], "Agoda Collect" if ln["collect_model"] == "OTA_COLLECT" else "Hotel Collect",
                   float(dot(ln["gross"])), float(dot(ln["base"])), float(ln["pct"]), float(dot(ln["commission"])),
                   CCY, doc_no])
    ws.append(["Total"] + [None] * 13 + [float(dot(sum(ln["commission"] for ln in lines)))])
    wb.properties.created = wb.properties.modified = FIXED_TS
    buf = io.BytesIO()
    wb.save(buf)
    return _stable_zip(buf.getvalue()), f"agoda-finance-{mkey(month)}.xlsx", XLSX_TYPE


def _stable_zip(data: bytes) -> bytes:
    """An XLSX is a zip whose entries carry the save time: rewrite them with
    a fixed timestamp so the same seed gives the same bytes (and sha256)."""
    import zipfile

    import re
    import zipfile

    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    stamp = FIXED_TS.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps the save time into dcterms:modified whatever
                # the workbook properties say.
                body = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rb"\g<1>" + stamp + rb"\g<2>",
                              body)
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 5, 8, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(fixed, body)
    return out.getvalue()


EXPEDIA_COLS = (("Reservation ID", 30), ("Guest", 44), ("Check-in", 24), ("Check-out", 24), ("Nights", 14),
                ("Status", 24), ("Booking amount", 30), ("Commissionable", 30), ("Comm. %", 18), ("Commission", 28))
EXPEDIA_STATUS = {"OK": "Postdeparture", "NO_SHOW": "No-show", "CANCELLED": "Cancelled"}


def _latin(s: Any) -> str:
    return str(s).encode("latin-1", "replace").decode("latin-1")


def render_expedia_pdf(lines: list[dict], *, doc_no: str, month: date, hotel_name: str) -> tuple[bytes, str, str]:
    """Expedia's Hotel Collect commission invoice for one month (fpdf2, core
    fonts), with a text layer pypdf reads back."""
    from fpdf import FPDF

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_creation_date(FIXED_TS.replace(tzinfo=UTC))
    pdf.set_title(f"Expedia commission invoice {doc_no}")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, _latin(f"Expedia Group - Hotel Collect commission invoice"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 9)
    for text in (f"Invoice number: {doc_no}", f"Property: {hotel_name}",
                 f"Period: {month.isoformat()} to {month_end(month).isoformat()}", "Currency: AED",
                 "Issued by: Expedia Lodging Partner Services Sarl"):
        pdf.cell(0, 5, _latin(text), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 8)
    for label, w in EXPEDIA_COLS:
        pdf.cell(w, 6, label, border=1)
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    for ln in lines:
        cells = (ln["ota_res_id"], ln["guest"], ln["check_in"].isoformat(), ln["check_out"].isoformat(),
                 ln["nights"], EXPEDIA_STATUS[ln["status"]], dot(ln["gross"]), dot(ln["base"]),
                 f"{ln['pct']:.2f}%", dot(ln["commission"]))
        for (_label, w), c in zip(EXPEDIA_COLS, cells):
            pdf.cell(w, 5, _latin(c), border=1)
        pdf.ln()
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 7, f"Total commission due: AED {dot(sum(ln['commission'] for ln in lines))}",
             new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output()), f"expedia-hc-invoice-{doc_no}.pdf", "application/pdf"


def expedia_extraction(lines: list[dict], *, doc_no: str, month: date, pdf_bytes: bytes) -> dict:
    """What the ingest.extract_pdf model returns for that PDF: one line per
    table row with the page it is printed on (read back from the PDF's text
    layer), values as printed, and the printed total."""
    from pypdf import PdfReader

    texts = [(p.extract_text() or "") for p in PdfReader(io.BytesIO(pdf_bytes)).pages]
    out = []
    for ln in lines:
        cells = {"ota_res_id": ln["ota_res_id"], "guest_name": ln["guest"], "check_in": ln["check_in"].isoformat(),
                 "check_out": ln["check_out"].isoformat(), "room_nights": str(ln["nights"]),
                 "line_status": EXPEDIA_STATUS[ln["status"]], "collect_model": "HOTEL_COLLECT",
                 "gross_minor": dot(ln["gross"]), "commissionable_minor": dot(ln["base"]),
                 "commission_pct": f"{ln['pct']:.2f}%", "commission_minor": dot(ln["commission"]), "currency": CCY}
        page = next((i for i, t in enumerate(texts, start=1) if ln["ota_res_id"] in t), 1)
        out.append({"page": page, "line_ref": ln["ota_res_id"],
                    "cells": [{"field": k, "value": v} for k, v in cells.items()]})
    return {"lines": out, "printed_total": dot(sum(ln["commission"] for ln in lines)), "currency": CCY,
            "doc_no": doc_no, "period_start": month.isoformat(), "period_end": month_end(month).isoformat()}


# ─── Import through clefs.files and core/ingest ───


def import_file(m: Money, key: str, rendered: tuple[bytes, str, str], *, kind: str, source_code: str | None = None,
                ota_code: str | None = None, extraction: dict | None = None) -> dict:
    """Store the bytes, stage them (sniff, alias mapping, printed-total
    gate) and commit through the kind's registered handler: exactly the
    console's upload path, with via=DEMO. `extraction` answers the one AI
    call a PDF makes (see the module docstring)."""
    from clefs import ai, files
    from clefs.core import ingest

    data, filename, content_type = rendered
    m.files[filename] = data
    owner = m.ctx.ids.get("users", {}).get("owner")
    ref = files.put_file(m.cur, m.slug, data, filename=filename, content_type=content_type, purpose="ingest.upload",
                         uploaded_by=owner)
    prev = getattr(ai, "_backend_override", None)
    ai_log = logging.getLogger("clefs.ai")
    level = ai_log.level
    if extraction is not None:
        ai.set_backend(ai.FakeBackend(extraction))
        # clefs.ai logs each call on its own connection; on a first seed the
        # hotel row is not committed yet, so that log row fails its foreign
        # key (caught by clefs.ai, which then prints a traceback). The replay
        # is no model call worth logging: keep the output quiet for it.
        ai_log.setLevel(logging.CRITICAL)
    try:
        batch = ingest.create_batch(m.cur, m.slug, file_id=ref.id, kind=kind, source_code=source_code,
                                    ota_code=ota_code, via="DEMO", uploaded_by=owner, use_ai=extraction is not None)
    finally:
        if extraction is not None:
            ai.set_backend(prev)
            ai_log.setLevel(level)
    if batch["status"] != "REVIEW" or batch.get("gate_blocking"):
        raise RuntimeError(f"money seed: {filename} did not stage cleanly: status={batch['status']} "
                           f"gate={batch.get('gate_status')} errors={batch.get('errors')} stats={batch.get('stats')}")
    counts = ingest.row_counts(m.cur, batch["id"])
    if counts.get("ERROR") or counts.get("DUP"):
        bad = [(r.get("row_no"), r.get("messages")) for r in ingest.list_rows(m.cur, batch["id"], limit=500)
               if r.get("status") in ("ERROR", "DUP")][:3]
        raise RuntimeError(f"money seed: {filename} has bad rows {counts}: {bad}")
    out = ingest.commit_batch(m.cur, m.slug, batch["id"])
    m.batches[key] = out
    m.bump("imports")
    return out


# ─── Detection: the registered job and task bodies ───
# Run inside the seed's transaction with a JobCtx, as the worker would run
# them (the worker commits after the body returns; here the seed does).


def run_task(m: Money, name: str, payload: dict) -> Any:
    from clefs import jobs

    t = next((t for t in jobs.registered_tasks() if t.name == name), None)
    if t is None:
        raise RuntimeError(f"money seed: task {name} is not registered")
    return t.fn(jobs.JobCtx(cur=m.cur, job=name, hotel_slug=m.slug, payload=payload), payload)


def run_job(m: Money, name: str) -> Any:
    from clefs import jobs

    j = next((j for j in jobs.registered_jobs() if j.name == name), None)
    if j is None:
        raise RuntimeError(f"money seed: job {name} is not registered")
    return j.fn(jobs.JobCtx(cur=m.cur, job=name, hotel_slug=m.slug, scheduled_for=clock.now()))


def detect(m: Money, batch_key: str | None = None) -> dict:
    """pg.detect_on_commit for one committed batch (its kind's detectors),
    or with no batch a full manual run of L1-L15."""
    if batch_key is None:
        return run_task(m, "pg.detect_on_commit", {"kind": "manual"})
    b = m.batches[batch_key]
    return run_task(m, "pg.detect_on_commit", {"kind": b["kind"], "batch_id": b["id"],
                                               "ota_code": (b.get("commit_result") or {}).get("ota_code")})


# ─── Plant stays for the statements ───

RATES = (275, 295, 310, 330, 355, 380)
L1_COUNT, L2_COUNT, L2_DECOYS, L4_COUNT, L5_COUNT, L6_COUNT = 25, 6, 5, 3, 3, 2
L5_PMS_RATE, L5_OTA_RATE = 420, 480
L6_PROGRAMS, L6_DECOY_PROGRAMS = "GENIUS, MOBILE, EARLY_BOOKER", "GENIUS, MOBILE"


def _rate(m: Money) -> int:
    return m.rng.choice(RATES)


def statement_stays(m: Money) -> None:
    """The stays behind every statement plant. August check-ins are after
    the 1 Aug fee configuration and after Preferred ended (31 Jul)."""
    for i in range(L1_COUNT):          # L1: commission on the full gross (MF + TD inside the base)
        make_stay(m, f"L1:{i}", channel="BOOKING", check_in=AUG + timedelta(days=1 + i % 24),
                  nights=2 + i % 3, rate=_rate(m), room_type=("STD_KING", "STD_TWIN")[i % 2])
    for i in range(L2_COUNT):          # L2: 18% in August, Preferred over
        make_stay(m, f"L2:{i}", channel="BOOKING", check_in=AUG + timedelta(days=2 + 4 * i), nights=3, rate=_rate(m))
    for i in range(L2_DECOYS):         # L2 decoy: 18% in July while Preferred was on
        make_stay(m, f"L2d:{i}", channel="BOOKING", check_in=JUL + timedelta(days=4 + 5 * i), nights=2, rate=_rate(m))
    # L3: no-shows the OTA still charged, a waived cancellation charged; the
    # decoy is a July no-show whose fee was collected (commission on the fee).
    make_stay(m, "L3:bk0", channel="BOOKING", check_in=AUG + timedelta(days=8), nights=2, rate=_rate(m),
              status="NO_SHOW", audit=False)
    make_stay(m, "L3:bk1", channel="BOOKING", check_in=AUG + timedelta(days=20), nights=2, rate=_rate(m),
              status="NO_SHOW", audit=False)
    make_stay(m, "L3:cxl", channel="BOOKING", check_in=AUG + timedelta(days=13), nights=3, rate=_rate(m),
              status="CANCELLED", cancel_fee=0, audit=False)
    make_stay(m, "L3:ag", channel="AGODA", check_in=AUG + timedelta(days=17), nights=2, rate=_rate(m),
              status="NO_SHOW", collect_model="OTA_COLLECT", audit=False)
    make_stay(m, "L3:ex", channel="EXPEDIA", check_in=AUG + timedelta(days=22), nights=2, rate=_rate(m),
              status="NO_SHOW", collect_model="HOTEL_COLLECT", audit=False)
    make_stay(m, "L3d", channel="BOOKING", check_in=JUL + timedelta(days=15), nights=2, rate=300,
              status="NO_SHOW", no_show_fee=30000, audit=False)
    for i in range(L4_COUNT):          # L4: 5 nights booked, 3 stayed, OTA not told
        make_stay(m, f"L4:{i}", channel="BOOKING", check_in=AUG + timedelta(days=3 + 7 * i), nights=5,
                  rate=_rate(m), actual_nights=3)
    make_stay(m, "L4d", channel="BOOKING", check_in=AUG + timedelta(days=24), nights=5, rate=_rate(m),
              actual_nights=3)         # decoy: reported correctly
    for i in range(L5_COUNT):          # L5: OTA still prices 480, PMS 420 after the modification
        make_stay(m, f"L5:{i}", channel="BOOKING", check_in=AUG + timedelta(days=5 + 6 * i), nights=2,
                  rate=L5_PMS_RATE)
    for i in range(L6_COUNT):          # L6: Genius 15 + Mobile 10 + Early Booker 15
        make_stay(m, f"L6:{i}", channel="BOOKING", check_in=AUG + timedelta(days=9 + 8 * i), nights=3,
                  rate=_rate(m), rate_plan="GENIUS_L2")
    make_stay(m, "L6d", channel="BOOKING", check_in=AUG + timedelta(days=11), nights=3, rate=_rate(m),
              rate_plan="GENIUS_L1")   # decoy: Genius 10 + Mobile 10 = 19%


def scope_genius(m: Money) -> None:
    """core_catalog seeds Genius 10 and Genius 15 as two GENIUS programs with
    no rate plans, so any line naming GENIUS would stack both levels. Genius
    is one level per booker: scope each row to its level's rate plan."""
    m.cur.execute(
        """
        UPDATE pfg_ota_programs SET rate_plans = CASE WHEN pct = 10 THEN ARRAY['GENIUS_L1']
                                                      ELSE ARRAY['GENIUS_L2', 'GENIUS_L3'] END, updated_at = NOW()
         WHERE hotel_slug = %s AND ota_code = 'BOOKING' AND program = 'GENIUS' AND kind = 'DISCOUNT'
        """,
        (m.slug,),
    )


# ─── Statement lines and their plants ───


def booking_doc(month: date) -> str:
    return f"BDC-{month:%Y%m}-{40000 + 17 * month.month}"


def _ids(*stays: dict) -> dict:
    """Stable references only (database ids change on every re-seed)."""
    return {"pms_res_ids": [s["pms_res_id"] for s in stays], "ota_res_ids": [s["ota_res_id"] for s in stays]}


def _one(m: Money, det: str, s: dict, *, ota: str, doc: str, amount: int, klass: str = "RECOVERABLE",
         note: str | None = None, extra: dict | None = None) -> None:
    plant(m, det, f"{det}:{ota}:{s['ota_res_id'] or s['pms_res_id']}", klass=klass, expected_minor=amount,
          fingerprints=[fingerprint(det, ota, s["norm"], doc)], ids=_ids(s) | (extra or {}), note=note)


def _decoys(m: Money, det: str, key: str, stays: list[dict], *, ota: str, doc: str, note: str) -> None:
    plant(m, det, f"{det}:decoys:{key}", klass="DECOY", expected_minor=0, fingerprints=[], ids={},
          decoy_ids=_ids(*stays), decoy_fingerprints=[fingerprint(det, ota, s["norm"], doc) for s in stays],
          note=note)


def booking_lines(m: Money) -> dict[date, list[dict]]:
    """Every Booking line of the three statements: core stays priced from
    the PMS, then the plants (and decoys) of July and August."""
    out = {month: core_lines(m, "BOOKING", month) for month in STATEMENT_MONTHS}
    jul, aug, doc_jul, doc_aug = out[JUL], out[AUG], booking_doc(JUL), booking_doc(AUG)
    decoys = [m.stays[f"L2d:{i}"] for i in range(L2_DECOYS)]
    for s in decoys:
        jul.append(make_line(stay_row(s), pms_view(m, "BOOKING", s["id"], s["check_in"])))
    _decoys(m, "L2", "preferred-active", decoys, ota="BOOKING", doc=doc_jul,
            note="18% on July stays while Preferred (+3) was active: due")
    s = m.stays["L3d"]
    pct = pms_view(m, "BOOKING", s["id"], s["check_in"])["pct"]
    fee = s["no_show_fee_minor"]
    jul.append(make_line(stay_row(s), {"gross": fee, "base": fee, "pct": pct, "commission": pct_of(fee, pct)},
                         status="NO_SHOW", remarks="No-show fee charged"))
    _decoys(m, "L3", "fee-collected", [s], ota="BOOKING", doc=f"M{mkey(s['check_out'])}",
            note="no-show fee collected, commission on the fee only")
    for i in range(L1_COUNT):
        s = m.stays[f"L1:{i}"]
        v = pms_view(m, "BOOKING", s["id"], s["check_in"])
        aug.append(make_line(stay_row(s), v, base=v["gross"], commission=pct_of(v["gross"], v["pct"])))
        _one(m, "L1", s, ota="BOOKING", doc=doc_aug, amount=pct_of(v["mf"] + v["td"], v["pct"]),
             note="commission on the full gross while MF and TD are configured as charges")
    for i in range(L2_COUNT):
        s = m.stays[f"L2:{i}"]
        v = pms_view(m, "BOOKING", s["id"], s["check_in"])
        comm = pct_of(v["base"], Decimal(18))
        aug.append(make_line(stay_row(s), v, pct=Decimal(18), commission=comm))
        _one(m, "L2", s, ota="BOOKING", doc=doc_aug, amount=comm - pct_of(v["base"], v["pct"]),
             note="18% after Preferred ended on 31 Jul")
    for key in ("L3:bk0", "L3:bk1", "L3:cxl"):
        s = m.stays[key]
        v = quoted_view(m, "BOOKING", s["check_in"], s["rates"])
        aug.append(make_line(stay_row(s), v, status="CANCELLED" if key == "L3:cxl" else "OK"))
        _one(m, "L3", s, ota="BOOKING", doc=f"M{mkey(s['check_out'])}", amount=v["commission"],
             note="waived cancellation charged" if key == "L3:cxl" else
             f"PMS no-show, commission charged; credited by {CREDIT_DOC}, so RECOVERED")
    for i in range(L4_COUNT):
        s = m.stays[f"L4:{i}"]
        v = quoted_view(m, "BOOKING", s["check_in"], s["rates"])
        aug.append(make_line(stay_row(s), v))
        _one(m, "L4", s, ota="BOOKING", doc=doc_aug, amount=rnd(Decimal(v["commission"]) * 2 / 5),
             note="5 nights booked, 3 stayed, OTA still reports 5")
    s = m.stays["L4d"]
    v = quoted_view(m, "BOOKING", s["check_in"], s["rates"][:3])
    aug.append(make_line(stay_row(s), v, check_out=s["check_in"] + timedelta(days=3), nights=3))
    _decoys(m, "L4", "reported", [s], ota="BOOKING", doc=doc_aug, note="early departure reported to the OTA")
    for i in range(L5_COUNT):
        s = m.stays[f"L5:{i}"]
        pms, ota = pms_view(m, "BOOKING", s["id"], s["check_in"]), quoted_view(m, "BOOKING", s["check_in"],
                                                                                [L5_OTA_RATE] * s["nights"])
        aug.append(make_line(stay_row(s), ota))
        _one(m, "L5", s, ota="BOOKING", doc=doc_aug, amount=pct_of(ota["base"] - pms["base"], pms["pct"]),
             note="OTA prices AED 480 a night, the PMS AED 420 after the modification")
    for key, programs in [(f"L6:{i}", L6_PROGRAMS) for i in range(L6_COUNT)] + [("L6d", L6_DECOY_PROGRAMS)]:
        s = m.stays[key]
        v = pms_view(m, "BOOKING", s["id"], s["check_in"])
        aug.append(make_line(stay_row(s), v, programs=programs))
        if key == "L6d":
            _decoys(m, "L6", "genius-mobile", [s], ota="BOOKING", doc=doc_aug, note="Genius 10 + Mobile 10 = 19%")
            continue
        stacked = 1 - Decimal("0.85") * Decimal("0.90") * Decimal("0.85")
        public = Decimal(v["gross"]) / s["nights"] / (1 - stacked)
        _one(m, "L6", s, ota="BOOKING", doc=doc_aug, klass="PREVENTABLE",
             amount=rnd((stacked - Decimal("0.25")) * public * s["nights"]),
             note="stacked depth 35.0% over the 25% cap")
    return out


AGODA_DOC, EXPEDIA_DOC = f"AGD-{AUG:%Y%m}-7731", f"EXP-HC-{AUG:%Y%m}-0912"


def _no_show_line(m: Money, key: str, ota: str, lines: list[dict]) -> None:
    """A PMS no-show the OTA invoiced as a completed stay (L3)."""
    s = m.stays[key]
    v = quoted_view(m, ota, s["check_in"], s["rates"])
    lines.append(make_line(stay_row(s), v))
    _one(m, "L3", s, ota=ota, doc=f"M{mkey(s['check_out'])}", amount=v["commission"],
         note="PMS no-show, commission charged")


def agoda_lines(m: Money) -> list[dict]:
    lines = core_lines(m, "AGODA", AUG)
    structural = list(lines)
    _no_show_line(m, "L3:ag", "AGODA", lines)
    # Agoda's commission base (the RSR) includes VAT and the municipality
    # fee by contract: Profit Guard stores that as STRUCTURAL (never a
    # claim, never signalled) on every Agoda line it can read. The no-show
    # line is left out: its L3 supersedes the structural record.
    plant(m, "L1", "L1:AGODA:structural", klass="STRUCTURAL",
          expected_minor=sum(pct_of(inferred_fees(m, ln, ("mf", "vat")), ln["pct"]) for ln in structural),
          fingerprints=[fingerprint("L1", "AGODA", core_stays.normalize_ota_res_id("AGODA", ln["ota_res_id"]),
                                    AGODA_DOC) for ln in structural],
          ids={"doc_no": AGODA_DOC, "lines": len(structural)},
          note="by design: commission on the tax-inclusive RSR (net ADR input, not a claim)")
    return lines


def expedia_lines(m: Money) -> list[dict]:
    """Hotel Collect stays only: Expedia invoices commission on those; its
    Expedia Collect stays are paid net on a virtual card instead."""
    lines = core_lines(m, "EXPEDIA", AUG, collect_model="HOTEL_COLLECT")
    _no_show_line(m, "L3:ex", "EXPEDIA", lines)
    return lines


# Arrival of each document (hotel-local day and hour of the upload).
ARRIVALS = {JUN: (date(2026, 7, 3), 1), JUL: (date(2026, 8, 3), 2), AUG: (date(2026, 9, 3), 3)}


def import_statements(m: Money) -> dict[str, list[dict]]:
    """The three Booking statements, the Agoda XLSX and the Expedia PDF, in
    the order they arrived. Detection runs after the June statement (the
    fees were not configured yet: L1b) and after the August one (its no-shows
    become L3 findings before the credit note that settles two of them)."""
    bk = booking_lines(m)
    for month, (day, variant) in ARRIVALS.items():
        key = f"booking:{mkey(month)}"
        m.lines[key] = bk[month]
        if not bk[month]:
            continue  # a tiny --scale can leave a month with no Booking stay: no document
        with at(m.ctx, day, 9, 30):
            import_file(m, key, render_booking(bk[month], variant=variant, doc_no=booking_doc(month)),
                        kind="OTA_STATEMENT", source_code="booking_res_statement", ota_code="BOOKING")
            if month in (JUN, AUG):
                detect(m, key)
    m.lines["agoda"] = ag = agoda_lines(m)
    with at(m.ctx, date(2026, 9, 5), 11):
        import_file(m, "agoda", render_agoda(ag, doc_no=AGODA_DOC, month=AUG), kind="OTA_STATEMENT",
                    source_code="agoda_finance", ota_code="AGODA")
    m.lines["expedia"] = ex = expedia_lines(m)
    pdf = render_expedia_pdf(ex, doc_no=EXPEDIA_DOC, month=AUG, hotel_name=m.ctx.hotel.name)
    with at(m.ctx, date(2026, 9, 6), 10):
        import_file(m, "expedia", pdf, kind="OTA_STATEMENT", ota_code="EXPEDIA",
                    extraction=expedia_extraction(ex, doc_no=EXPEDIA_DOC, month=AUG, pdf_bytes=pdf[0]))
    if bk[JUN]:
        l1b_plant(m, bk[JUN])
    return bk


CREDIT_DOC = f"BDC-CN-{AUG:%Y%m}-0033"


def import_credit_note(m: Money) -> None:
    """Booking credits the two August no-show commissions it was asked
    about; auto-verification then marks those L3 findings RECOVERED."""
    aug = {ln["ota_res_id"]: ln for ln in m.lines[f"booking:{mkey(AUG)}"]}
    stays = [m.stays["L3:bk0"], m.stays["L3:bk1"]]
    head = ["Credit note number", "Invoice number", "Reservation number", "Guest name", "Check in", "Amount",
            "Currency", "Reason"]
    rows = [head] + [[CREDIT_DOC, booking_doc(AUG), s["ota_res_id"], s["guest_name"], s["check_in"].isoformat(),
                      dot(aug[s["ota_res_id"]]["commission"]), CCY, "No-show commission reversed"] for s in stays]
    total = sum(aug[s["ota_res_id"]]["commission"] for s in stays)
    rows.append(["Total", "", "", "", "", dot(total), "", ""])
    with at(m.ctx, date(2026, 9, 12), 14):
        import_file(m, "credit_note", (_csv(rows).encode("utf-8"), f"booking-credit-note-{CREDIT_DOC}.csv",
                                       "text/csv"), kind="CREDIT_NOTE", ota_code="BOOKING")
        run_task(m, "pg.auto_verify", {"ota_code": "BOOKING"})
    m.lines["credit_note"] = [{"ota_res_id": s["ota_res_id"], "commission": -aug[s["ota_res_id"]]["commission"],
                               "check_in": s["check_in"]} for s in stays]


def l15_plants(m: Money) -> None:
    """Reverse-charge VAT: 5% of each month's commission from a non-UAE
    entity (Booking NL, Agoda SG, Expedia CH), credit notes netted."""
    months: dict[tuple[str, str], int] = {}
    for key, ota in [(f"booking:{mkey(mo)}", "BOOKING") for mo in STATEMENT_MONTHS] + [("agoda", "AGODA"),
                                                                                      ("expedia", "EXPEDIA")]:
        lines = m.lines.get(key) or []
        if not lines:
            continue
        month = mkey(max(ln["check_out"] for ln in lines))
        months[(ota, month)] = months.get((ota, month), 0) + sum(ln["commission"] for ln in lines)
    months[("BOOKING", mkey(AUG))] += sum(ln["commission"] for ln in m.lines.get("credit_note", []))
    for (ota, month), commission in sorted(months.items()):
        plant(m, "L15", f"L15:{ota}:{month}", klass="COMPLIANCE", expected_minor=pct_of(commission, Decimal(5)),
              fingerprints=[fingerprint("L15", ota, None, month)], ids={"ota": ota, "month": month},
              note="self-account 5% VAT on commission invoiced from abroad")


# ─── Payouts (Payments by Booking.com) ───

PAY_FEE_PCT = Decimal("1.2")          # payments service fee on the amount the guest paid
CHARGEBACK_AGES = (20, 55, 95)
L10_SHORT = (8500, 14200)             # two lines paid short
L10_BANK_SHORT = 124000               # one bank credit AED 1,240 short
REFUND_MINOR = 60000                  # two partial refunds of AED 600
USD_BAD_RATE, EUR_SPREAD = Decimal("3.61"), Decimal("0.025")
EUR_AED_REF = Decimal("4.2915")       # reference rate written only when fx_rates has no EUR/AED row


def payout_stays(m: Money) -> None:
    """OTA-collect Booking stays behind the payout plants (L10-L13)."""
    a = m.ctx.anchor
    for i in range(2):
        make_stay(m, f"L10:{i}", channel="BOOKING", check_in=a - timedelta(days=22 - 5 * i), nights=2, rate=_rate(m),
                  collect_model="OTA_COLLECT")
    # Large enough on its own that the credit, AED 1,265 under it (the short
    # plus the fee), is still a credit whatever else shares its payout.
    make_stay(m, "L10:bank", channel="BOOKING", check_in=a - timedelta(days=30), nights=5, rate=380,
              collect_model="OTA_COLLECT")
    for i in range(2):
        make_stay(m, f"L12:{i}", channel="BOOKING", check_in=AUG + timedelta(days=12 + 6 * i), nights=4,
                  rate=_rate(m), collect_model="OTA_COLLECT")
    for i, age in enumerate(CHARGEBACK_AGES):
        make_stay(m, f"L13:{i}", channel="BOOKING", check_in=a - timedelta(days=age + 9), nights=3, rate=_rate(m),
                  collect_model="OTA_COLLECT")
    for key, back in (("L11:usd", 26), ("L11:eur", 19), ("L11d", 12)):
        make_stay(m, key, channel="BOOKING", check_in=a - timedelta(days=back), nights=2, rate=_rate(m),
                  collect_model="OTA_COLLECT")


def _payday(d: date) -> date:
    """Booking pays out weekly, on the Thursday after the event."""
    return d + timedelta(days=(3 - d.weekday()) % 7 or 7)


def _pline(r: dict, v: dict, *, kind: str = "Reservation", event: date | None = None) -> dict:
    fee = pct_of(v["gross"], PAY_FEE_PCT)
    return {"reservation_number": r["ota_res_id"], "guest_name": r["guest_name"], "check_in": r["check_in"],
            "check_out": r["check_out"], "type": kind, "gross_amount": v["gross"], "commissionable_amount": v["base"],
            "commission": v["commission"], "payments_service_fee": fee,
            "net_amount": v["gross"] - v["commission"] - fee, "event_date": event or r["check_out"]}


def _convert(ln: dict, rate: Decimal) -> dict:
    """A line re-expressed in a payout currency (AED / rate)."""
    out = dict(ln)
    for k in ("gross_amount", "commissionable_amount", "commission", "payments_service_fee"):
        out[k] = rnd(Decimal(ln[k]) / rate)
    out["net_amount"] = out["gross_amount"] - out["commission"] - out["payments_service_fee"]
    return out


def eur_rate(m: Money, on: date) -> Decimal:
    """EUR/AED from fx_rates; written once (MANUAL, a plausible reference
    rate) only when the table has no EUR/AED row at all on or before `on`,
    because L11 needs a reference rate to measure the spread against."""
    from clefs.core import money as core_money

    rate = core_money.fx_rate(m.cur, "EUR", "AED", on)
    if rate is None:
        m.cur.execute("INSERT INTO fx_rates (date, base, quote, rate, source) VALUES (%s, 'EUR', 'AED', %s, 'MANUAL') "
                      "ON CONFLICT DO NOTHING", (on - timedelta(days=1), EUR_AED_REF))
        rate = core_money.fx_rate(m.cur, "EUR", "AED", on)
    return Decimal(rate)


def build_payouts(m: Money) -> list[dict]:
    """Weekly AED payouts of every OTA-collect Booking stay (core and plant),
    then one USD, one EUR and one USD decoy payout. Each payout: {ref, date,
    currency, bank_ref, lines, plants}."""
    a, s = m.ctx.anchor, m.stays
    cutoff = a - timedelta(days=2)
    m.cur.execute(_CORE_SQL, (m.slug, "BOOKING", JUN, cutoff - timedelta(days=1), "OTA_COLLECT", "OTA_COLLECT"))
    events = [_pline(r, pms_view(m, "BOOKING", r["id"], r["check_in"])) for r in [dict(x) for x in m.cur.fetchall()]]
    tags: dict[int, str] = {}
    for key in [f"L10:{i}" for i in range(2)] + ["L10:bank"] + [f"L12:{i}" for i in range(2)] + \
               [f"L13:{i}" for i in range(3)]:
        st = s[key]
        ln = _pline(stay_row(st) | {"guest_name": st["guest_name"]}, pms_view(m, "BOOKING", st["id"], st["check_in"]))
        if key.startswith("L10:") and key != "L10:bank":
            ln["net_amount"] -= L10_SHORT[int(key[-1])]
        events.append(ln)
        tags[id(ln)] = key
        if key.startswith("L12:"):
            ref = {"type": "Refund", "gross_amount": -REFUND_MINOR, "commissionable_amount": None, "commission": 0,
                   "payments_service_fee": 0, "net_amount": -REFUND_MINOR,
                   "event_date": st["check_out"] + timedelta(days=9)}
            events.append(ln | ref)
            tags[id(events[-1])] = key + ":refund"
        if key.startswith("L13:"):
            cb = {"type": "Chargeback", "gross_amount": None, "commissionable_amount": None, "commission": None,
                  "payments_service_fee": None, "net_amount": -ln["gross_amount"],
                  "event_date": a - timedelta(days=CHARGEBACK_AGES[int(key[-1])])}
            events.append(ln | cb)
            tags[id(events[-1])] = key + ":cb"
    payouts: dict[date, dict] = {}
    for ln in events:
        day = _payday(ln["event_date"])
        if day > cutoff:
            continue
        p = payouts.setdefault(day, {"ref": f"PO-{day:%Y%m%d}", "date": day, "currency": CCY,
                                     "bank_ref": f"BDCPAY{day:%y%m%d}AE", "lines": [], "plants": []})
        p["lines"].append(ln)
        if id(ln) in tags:
            p["plants"].append(tags[id(ln)])
    out = [payouts[d] for d in sorted(payouts)]
    usd = Decimal("3.6725")
    for key, ccy in (("L11:usd", "USD"), ("L11:eur", "EUR"), ("L11d", "USD")):
        st = s[key]
        day = _payday(st["check_out"])
        rate = usd if ccy == "USD" else eur_rate(m, day + timedelta(days=1))
        ln = _convert(_pline(stay_row(st) | {"guest_name": st["guest_name"]},
                             pms_view(m, "BOOKING", st["id"], st["check_in"])), rate)
        # The currency leads the reference: the bank linker matches a payout
        # reference inside the credit's text, and "PO-<date>" (the AED payout
        # of the same day) must not be a substring of this one.
        out.append({"ref": f"PO-{ccy}-{day:%Y%m%d}-{st['ota_res_id'][-4:]}", "date": day, "currency": ccy,
                    "bank_ref": f"BDCPAY{ccy}{day:%y%m%d}{st['ota_res_id'][-4:]}", "lines": [ln], "plants": [key],
                    "rate": rate})
    return out


PAYOUT_HEADERS = ("Payout ID", "Payout date", "Reservation number", "Guest name", "Check-in", "Check-out", "Type",
                  "Gross amount", "Commissionable amount", "Commission", "Payments service fee", "Net amount",
                  "Currency", "Bank transfer reference", "Event date")


def _amt(v: int | None) -> str | None:
    return None if v is None else dot(v)


def render_payouts_json(payouts: list[dict]) -> bytes:
    """The Payouts API shape (the raw artifact kept in files/)."""
    body = {"payouts": [{
        "payout_id": p["ref"], "payout_date": p["date"].isoformat(), "currency": p["currency"],
        "amount": dot(sum(ln["net_amount"] for ln in p["lines"])), "bank_transfer_reference": p["bank_ref"],
        "transactions": [{"reservation_number": ln["reservation_number"], "guest_name": ln["guest_name"],
                          "check_in": ln["check_in"].isoformat(), "check_out": ln["check_out"].isoformat(),
                          "type": ln["type"], "gross_amount": _amt(ln["gross_amount"]),
                          "commissionable_amount": _amt(ln["commissionable_amount"]),
                          "commission": _amt(ln["commission"]),
                          "payments_service_fee": _amt(ln["payments_service_fee"]),
                          "net_amount": dot(ln["net_amount"]), "event_date": ln["event_date"].isoformat()}
                         for ln in p["lines"]]} for p in payouts]}
    return (json.dumps(body, indent=2, sort_keys=True) + "\n").encode("utf-8")


def render_payouts_csv(payouts: list[dict]) -> bytes:
    rows: list[list[Any]] = [list(PAYOUT_HEADERS)]
    for p in payouts:
        for ln in p["lines"]:
            rows.append([p["ref"], p["date"].isoformat(), ln["reservation_number"], ln["guest_name"],
                         ln["check_in"].isoformat(), ln["check_out"].isoformat(), ln["type"],
                         _amt(ln["gross_amount"]), _amt(ln["commissionable_amount"]), _amt(ln["commission"]),
                         _amt(ln["payments_service_fee"]), dot(ln["net_amount"]), p["currency"], p["bank_ref"],
                         ln["event_date"].isoformat()])
    return _csv(rows).encode("utf-8")


def bank_credits(payouts: list[dict]) -> list[dict]:
    """What reached the hotel's AED account: every payout less the bank's
    AED 25 fee (L10's expected fee), one of them AED 1,240 short, and the
    foreign payouts converted at the bank's rate."""
    out = []
    for p in payouts:
        total = sum(ln["net_amount"] for ln in p["lines"])
        if total <= 0:
            continue
        if p["currency"] == CCY:
            amount, note = total - BANK_FEE_MINOR, f"Booking.com payout {p['ref']}"
            if "L10:bank" in p["plants"]:
                amount -= L10_BANK_SHORT
        else:
            rate = {"L11:usd": USD_BAD_RATE, "L11:eur": p["rate"] * (1 - EUR_SPREAD)}.get(p["plants"][0], p["rate"])
            amount, note = rnd(Decimal(total) * rate), f"FX {p['currency']} {dot(total)} @ {rate:.4f} payout {p['ref']}"
        out.append({"value_date": p["date"] + timedelta(days=1), "amount": amount, "reference": p["bank_ref"],
                    "counterparty": "Booking.com B.V.", "description": note, "payout": p})
    return out


BANK_DEBITS = ((date(2026, 7, 8), -845000, "DEWA-0726", "Dubai Electricity and Water Authority", "Utilities"),
               (date(2026, 8, 20), -1230000, "RENT-0826", "Al Rigga Properties LLC", "Lease instalment"),
               (date(2026, 9, 3), -312000, "SUP-4471", "Gulf Linen Supplies LLC", "Laundry"))


def render_bank(credits: list[dict]) -> bytes:
    rows: list[list[Any]] = [["Value date", "Amount", "Currency", "Reference", "Counterparty", "Description"]]
    items = [(c["value_date"], c["amount"], c["reference"], c["counterparty"], c["description"]) for c in credits]
    items += [d for d in BANK_DEBITS]
    for day, amount, ref, who, what in sorted(items, key=lambda x: (x[0], x[2])):
        rows.append([day.strftime("%d/%m/%Y"), dot(amount), CCY, ref, who, what])
    return _csv(rows).encode("utf-8")


def _payout_of(payouts: list[dict], tag: str) -> dict:
    return next(p for p in payouts if tag in p["plants"])


def import_payouts(m: Money) -> None:
    """Payouts JSON (kept) -> CSV export (imported), then the bank CSV, and
    the L10-L13 plants with the payout references they landed in."""
    m.cur.execute(
        """
        INSERT INTO pfg_detector_tuning (hotel_slug, detector, params) VALUES (%s, 'L10', %s)
        ON CONFLICT (hotel_slug, detector) DO UPDATE SET params = pfg_detector_tuning.params || EXCLUDED.params
        """,
        (m.slug, json.dumps({"bank_fee_minor": BANK_FEE_MINOR})),
    )
    payouts = build_payouts(m)
    a = m.ctx.anchor
    m.files["booking-payouts.json"] = render_payouts_json(payouts)
    with at(m.ctx, a - timedelta(days=2), 10):
        import_file(m, "payouts", (render_payouts_csv(payouts), "booking-payouts.csv", "text/csv"),
                    kind="OTA_PAYOUT", source_code="booking_payouts", ota_code="BOOKING")
    credits = bank_credits(payouts)
    with at(m.ctx, a - timedelta(days=1), 16):
        import_file(m, "bank", (render_bank(credits), "bank-statement-aed.csv", "text/csv"), kind="BANK")
    s = m.stays
    for i in range(2):
        st = s[f"L10:{i}"]
        _one(m, "L10", st, ota="BOOKING", doc=_payout_of(payouts, f"L10:{i}")["ref"], amount=L10_SHORT[i],
             note="payout line below gross - commission - fee")
    p = _payout_of(payouts, "L10:bank")
    plant(m, "L10", f"L10:bank:{p['ref']}", klass="RECOVERABLE", expected_minor=L10_BANK_SHORT,
          fingerprints=[fingerprint("L10", "BOOKING", "bank", p["ref"])], ids={"payout_ref": p["ref"]},
          note="bank credit AED 1,240 below the payout less the AED 25 fee")
    clean = [p for p in payouts if p["currency"] == CCY and not p["plants"]]
    if clean:
        plant(m, "L10", "L10:decoys:bank-fee", klass="DECOY", expected_minor=0, fingerprints=[], ids={},
              decoy_ids={"payout_ref": clean[-1]["ref"]},
              decoy_fingerprints=[fingerprint("L10", "BOOKING", "bank", clean[-1]["ref"])],
              note="credit = payout - the expected AED 25 bank fee")
    for key in ("L11:usd", "L11:eur"):
        p = _payout_of(payouts, key)
        c = next(c for c in credits if c["payout"] is p)
        total = sum(ln["net_amount"] for ln in p["lines"])
        plant(m, "L11", f"L11:{p['ref']}", klass="PREVENTABLE", currency=CCY,
              expected_minor=rnd(Decimal(total) * p["rate"]) - c["amount"],
              fingerprints=[fingerprint("L11", "BOOKING", "fx", p["ref"])], ids={"payout_ref": p["ref"]},
              note="USD converted at 3.61 (peg 3.6725)" if key == "L11:usd" else "EUR converted with a 2.5% spread")
    p = _payout_of(payouts, "L11d")
    plant(m, "L11", "L11:decoys:peg", klass="DECOY", expected_minor=0, fingerprints=[], ids={},
          decoy_ids={"payout_ref": p["ref"]}, decoy_fingerprints=[fingerprint("L11", "BOOKING", "fx", p["ref"])],
          note="USD converted at the 3.6725 peg")
    for i in range(2):
        st = s[f"L12:{i}"]
        v = pms_view(m, "BOOKING", st["id"], st["check_in"])
        _one(m, "L12", st, ota="BOOKING", doc=_payout_of(payouts, f"L12:{i}:refund")["ref"],
             amount=rnd(Decimal(REFUND_MINOR) * v["commission"] / v["base"]),
             note="AED 600 refunded, commission not reduced")
    for i in range(3):
        st = s[f"L13:{i}"]
        v = pms_view(m, "BOOKING", st["id"], st["check_in"])
        _one(m, "L13", st, ota="BOOKING", doc=f"CB:{_payout_of(payouts, f'L13:{i}:cb')['ref']}",
             amount=v["gross"], note=f"chargeback aged {CHARGEBACK_AGES[i]} days")


def inferred_fees(m: Money, ln: dict, comps: tuple[str, ...] = ("mf", "td", "tf", "dest")) -> int:
    """The fee components inside a printed gross, per night (default MF +
    TD + TF + destination fee: what Booking's excluded charges cover)."""
    nights = max(1, ln["nights"] * ln["rooms"])  # the printed room nights, as the detector reads them
    per, rest = divmod(int(ln["gross"]), nights)
    total = 0
    for i in range(nights):
        night = core_tax.decompose_night(m.cur, m.slug, gross_minor=per + (rest if i == nights - 1 else 0),
                                         stay_date=ln["check_in"] + timedelta(days=i), rooms=max(1, ln["rooms"]),
                                         td_liable=ln["status"] != "NO_SHOW")
        total += sum(int(night.get(c) or 0) for c in comps)
    return total


def l1b_plant(m: Money, june: list[dict]) -> None:
    """L1b, raised by the detection run after the June statement: the fees
    were still inside the rate, so 90 days of MF + TD commission / 3."""
    # The fees are read back out of each line's gross night by night, as the
    # detector reads them (core.tax.decompose_night on gross / nights, no
    # room type): the statement carries no price breakdown.
    total = sum(pct_of(inferred_fees(m, ln), ln["pct"]) for ln in june if ln["commission"] > 0)
    month = mkey(ARRIVALS[JUN][0])
    plant(m, "L1b", f"L1b:BOOKING:{month}", klass="PREVENTABLE", expected_minor=rnd(Decimal(total) / 3),
          fingerprints=[fingerprint("L1b", "BOOKING", None, month)],
          ids={"doc_no": booking_doc(JUN), "lines": len(june)},
          note="MF and TD not configured as charges before 1 Aug: commission on them each month")


# ─── Virtual cards ───

EVC_DAYS = (5, 5, 12, 12, 20, 40, 40)     # Expedia EVCs expiring this many days after the anchor
BOOKING_VCC_DAYS = (9, 25, 33)
LEGACY_EVCS = ((date(2026, 3, 9), 12, 84500), (date(2026, 3, 21), 4, 112000))  # activation, days expired, balance


def card_stays(m: Money) -> None:
    """Expedia Collect and Booking OTA-collect stays whose cards are still
    loaded (L8), and one Expedia stay whose card was fully charged (decoy)."""
    a = m.ctx.anchor
    for i in range(len(EVC_DAYS)):
        make_stay(m, f"L8:ex{i}", channel="EXPEDIA", check_in=a - timedelta(days=12 - i), nights=2, rate=_rate(m),
                  collect_model="OTA_COLLECT")
    for i in range(len(BOOKING_VCC_DAYS)):
        make_stay(m, f"L8:bk{i}", channel="BOOKING", check_in=a - timedelta(days=9 - i), nights=2, rate=_rate(m),
                  collect_model="OTA_COLLECT")
    make_stay(m, "L8d", channel="EXPEDIA", check_in=a - timedelta(days=8), nights=2, rate=_rate(m),
              collect_model="OTA_COLLECT")


VCC_HEADERS = ("Card ID", "Reservation number", "Activation date", "Expiration date", "Initial amount",
               "Current balance", "Status", "Currency")


def _card(vcc_id: str, ota_res_id: str, activation: date, expiry: date, initial: int, balance: int,
          status: str) -> list[Any]:
    return [vcc_id, ota_res_id, activation.isoformat(), expiry.isoformat(), dot(initial), dot(balance), status, CCY]


def _loaded(m: Money, key: str, ota: str, prefix: str, days: int, status: str, report: list) -> None:
    st = m.stays[key]
    v = pms_view(m, ota, st["id"], st["check_in"])
    vcc, net = f"{prefix}-{st['ota_res_id'][-7:]}", v["gross"] - v["commission"]
    report.append(_card(vcc, st["ota_res_id"], st["check_in"], m.ctx.anchor + timedelta(days=days), net, net,
                        status))
    _one(m, "L8", st, ota=ota, doc=f"VCC:{vcc}", amount=net, extra={"vcc_id": vcc},
         note=f"card loaded, charge deadline in {days} days")


def import_cards(m: Money) -> None:
    """The Expedia EVC report and the Booking VCC report as downloaded on
    the anchor morning: loaded cards (L8), two legacy EVCs past their
    deadline (L9), and fully charged cards (decoys, core and plant)."""
    a, s = m.ctx.anchor, m.stays
    reports: dict[str, list[list[Any]]] = {"EXPEDIA": [], "BOOKING": []}
    for ota in reports:
        m.cur.execute(_CORE_SQL, (m.slug, ota, a - timedelta(days=30), a, "OTA_COLLECT", "OTA_COLLECT"))
        for r in [dict(x) for x in m.cur.fetchall()]:
            v = pms_view(m, ota, r["id"], r["check_in"])
            vcc = f"{'EVC' if ota == 'EXPEDIA' else 'VCC'}-{r['ota_res_id'][-7:]}"
            reports[ota].append(_card(vcc, r["ota_res_id"], r["check_in"], r["check_out"] + timedelta(days=60),
                                      v["gross"] - v["commission"], 0, "Fully charged"))
    for i, days in enumerate(EVC_DAYS):
        _loaded(m, f"L8:ex{i}", "EXPEDIA", "EVC", days, "Active", reports["EXPEDIA"])
    for i, days in enumerate(BOOKING_VCC_DAYS):
        _loaded(m, f"L8:bk{i}", "BOOKING", "VCC", days, "Available", reports["BOOKING"])
    st = s["L8d"]
    v = pms_view(m, "EXPEDIA", st["id"], st["check_in"])
    vcc_d = f"EVC-{st['ota_res_id'][-7:]}"
    reports["EXPEDIA"].append(_card(vcc_d, st["ota_res_id"], st["check_in"], a + timedelta(days=15),
                                    v["gross"] - v["commission"], 0, "Fully charged"))
    plant(m, "L8", "L8:decoys:fully-charged", klass="DECOY", expected_minor=0, fingerprints=[], ids={},
          decoy_ids=_ids(st) | {"vcc_id": vcc_d},
          decoy_fingerprints=[fingerprint(d, "EXPEDIA", st["norm"], f"VCC:{vcc_d}") for d in ("L8", "L9")],
          note="card fully charged: nothing to collect")
    for activated, expired, balance in LEGACY_EVCS:
        ota_res_id = _ota_id(m, "EXPEDIA")
        vcc = f"EVC-{ota_res_id[-7:]}"
        reports["EXPEDIA"].append(_card(vcc, ota_res_id, activated, a - timedelta(days=expired), balance, balance,
                                        "Active"))
        plant(m, "L9", f"L9:EXPEDIA:{vcc}", klass="RECOVERABLE", expected_minor=balance,
              fingerprints=[fingerprint("L9", "EXPEDIA", core_stays.normalize_ota_res_id("EXPEDIA", ota_res_id),
                                        f"VCC:{vcc}")],
              ids={"vcc_id": vcc, "ota_res_ids": [ota_res_id]},
              note=f"legacy EVC from a March check-out, expired {expired} days ago with money on it")
    with at(m.ctx, a, 8, 15):
        for ota, rows in reports.items():
            total = sum(int(Decimal(r[5]) * 100) for r in rows)
            body = [list(VCC_HEADERS)] + rows + [["Total", "", "", "", "", dot(total), "", ""]]
            import_file(m, f"vcc:{ota.lower()}", (_csv(body).encode("utf-8"), f"{ota.lower()}-vcc-report.csv",
                                                  "text/csv"), kind="VCC_REPORT", ota_code=ota)


# ─── Deadlines before any statement (L14) ───


def l14_stays(m: Money) -> None:
    """Two Booking no-shows on the anchor's previous night (48 h window
    closes the day after tomorrow) and an Expedia Hotel Collect early
    departure (shortened stay) the OTA must hear about by the 4th."""
    a = m.ctx.anchor
    for i in range(2):
        make_stay(m, f"L14:bk{i}", channel="BOOKING", check_in=a - timedelta(days=1), nights=1, rate=_rate(m),
                  status="NO_SHOW", audit=False)
    make_stay(m, "L14:ex", channel="EXPEDIA", check_in=a - timedelta(days=12), nights=4, rate=_rate(m),
              actual_nights=2, collect_model="HOTEL_COLLECT")


def l14_plants(m: Money) -> None:
    for i in range(2):
        st = m.stays[f"L14:bk{i}"]
        v = pms_view(m, "BOOKING", st["id"], st["check_in"])
        _one(m, "L14", st, ota="BOOKING", doc=mkey(st["check_out"]), amount=v["commission"],
             note="PMS no-show: mark it in the extranet within 48 h of check-out")
    st = m.stays["L14:ex"]
    v = quoted_view(m, "EXPEDIA", st["check_in"] + timedelta(days=2), st["rates"][2:])
    _one(m, "L14", st, ota="EXPEDIA", doc=mkey(st["check_out"]), amount=v["commission"],
         note="Hotel Collect stay shortened: report it by the 4th, 12:00")


# ─── Parity evidence (L7) ───

L7_GAP = Decimal("0.06")         # Agoda 6% under the direct price
GENIUS_GAP = Decimal("0.10")     # the decoy: a Genius (logged-in) price 10% lower


def placeholder_png(label: str) -> bytes:
    """A small valid PNG standing in for a rate-shop screenshot, with the
    capture described in a tEXt chunk (distinct bytes per capture)."""
    import struct
    import zlib

    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)

    w, h = 64, 36
    raw = b"".join(b"\x00" + bytes((200, 200, 205)) * w for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"tEXt", b"Comment\x00" + label.encode("latin-1", "replace"))
            + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))


def _agoda_rooms(m: Money, d: date) -> int:
    m.cur.execute(
        """
        SELECT COALESCE(SUM(n.rooms), 0) AS rooms FROM reservation_nights n JOIN reservations r ON r.id = n.reservation_id
         WHERE n.hotel_slug = %s AND n.stay_date = %s AND UPPER(r.channel_code) = 'AGODA'
           AND r.status NOT IN ('CANCELLED', 'NO_SHOW')
        """,
        (m.slug, d),
    )
    return int(m.cur.fetchone()["rooms"])


def _evidence(m: Money, channel: str, d: date, price: int, *, login: str = "LOGGED_OUT", shot: str | None = None) -> str:
    """Record one capture through profit_guard.api; returns its comparable key."""
    ev = {"channel": channel, "stay_date": d, "price_minor": price, "currency": CCY, "los": 1, "adults": 2,
          "room_label": "Standard King", "plan_label": "Room only", "taxes_included": True, "device": "DESKTOP",
          "login_state": login, "booker_country": "AE", "captured_at": m.ctx.at(m.ctx.anchor - timedelta(days=1), 10),
          "screenshot_file_id": shot,
          "search_url": f"https://www.example.com/rate-shop/{channel.lower()}/{d.isoformat()}"}
    eid = m.api.record_rate_evidence(m.cur, m.slug, evidence=ev, emit_signal=True)
    if not eid:
        raise RuntimeError(f"money seed: rate evidence for {channel} {d} was refused")
    m.cur.execute("SELECT comparable_key FROM pfg_rate_evidence WHERE id = %s", (eid,))
    m.bump("rate_evidence")
    return m.cur.fetchone()["comparable_key"]


def rate_evidence(m: Money) -> None:
    """Three future nights where Agoda (logged out, desktop) sells 6% under
    the direct price, captured yesterday with screenshots; one night where
    only a Genius logged-in price is lower (the decoy: not comparable)."""
    from clefs import files

    a = m.ctx.anchor
    m.cur.execute(
        """
        SELECT n.stay_date, SUM(n.rooms) AS rooms FROM reservation_nights n JOIN reservations r ON r.id = n.reservation_id
         WHERE n.hotel_slug = %s AND n.stay_date BETWEEN %s AND %s AND UPPER(r.channel_code) = 'AGODA'
           AND r.status NOT IN ('CANCELLED', 'NO_SHOW')
         GROUP BY 1 ORDER BY 2 DESC, 1 LIMIT 3
        """,
        (m.slug, a + timedelta(days=3), a + timedelta(days=40)),
    )
    days = sorted(r["stay_date"] for r in m.cur.fetchall())
    for extra in (10, 17, 24):
        if len(days) < 3 and a + timedelta(days=extra) not in days:
            days.append(a + timedelta(days=extra))
    for i, d in enumerate(sorted(days)[:3]):
        direct = (320 + 15 * i) * 100
        agoda = rnd(Decimal(direct) * (1 - L7_GAP))
        png = placeholder_png(f"Agoda {d.isoformat()} AED {dot(agoda)} logged out desktop (DEMO placeholder)")
        shot = files.put_file(m.cur, m.slug, png, filename=f"agoda-{d.isoformat()}.png", content_type="image/png",
                              purpose="profit.rate_evidence").id
        _evidence(m, "DIRECT_WEB", d, direct)
        _evidence(m, "BOOKING", d, direct)
        key = _evidence(m, "AGODA", d, agoda, shot=shot)
        plant(m, "L7", f"L7:AGODA:{d.isoformat()}", klass="PREVENTABLE",
              expected_minor=(direct - agoda) * _agoda_rooms(m, d),
              fingerprints=[fingerprint("L7", "AGODA", key, d.isoformat())],
              ids={"stay_date": d.isoformat(), "comparable_key": key}, note="Agoda 6% under direct, screenshot kept")
    d = next(a + timedelta(days=k) for k in range(14, 60) if a + timedelta(days=k) not in days)
    direct = 34000
    _evidence(m, "DIRECT_WEB", d, direct)
    _evidence(m, "BOOKING", d, direct)
    key = _evidence(m, "BOOKING", d, rnd(Decimal(direct) * (1 - GENIUS_GAP)), login="LOGGED_IN")
    plant(m, "L7", "L7:decoys:genius", klass="DECOY", expected_minor=0, fingerprints=[], ids={},
          decoy_ids={"stay_date": d.isoformat(), "comparable_key": key},
          decoy_fingerprints=[fingerprint("L7", "BOOKING", key, d.isoformat())],
          note="Genius members' logged-in price: not comparable with a public price")


# ─── Tourism Dirham: folio anomalies (T1-T6), the ledger (T7), the exemption (TX1) ───

T3_POSTED = 1500                 # TD keyed as AED 15 instead of the per-unit rate


def tax_stays(m: Money) -> None:
    """Non-OTA stays whose folio (from the PMS export) carries the T plants."""
    make_stay(m, "T1", channel="CORPORATE", check_in=AUG + timedelta(days=10), nights=4, rate=260, rooms=3,
              rate_plan="CORP", audit=False)                      # 3 rooms, TD posted once a night
    make_stay(m, "T2:a", channel="DIRECT_WEB", check_in=AUG + timedelta(days=1), nights=25, rate=240, guest=5,
              rate_plan="BAR_DIRECT", audit=False)                # 25 + 20 nights back to back
    make_stay(m, "T2:b", channel="DIRECT_WEB", check_in=AUG + timedelta(days=26), nights=20, rate=240, guest=5,
              rate_plan="BAR_DIRECT", audit=False)
    make_stay(m, "T3", channel="DIRECT_WEB", check_in=AUG + timedelta(days=12), nights=10, rate=250,
              rate_plan="BAR_DIRECT", audit=False)                # TD keyed at AED 15
    make_stay(m, "T4", channel="WALKIN", check_in=AUG + timedelta(days=2), nights=7, rate=230, rate_plan="RACK",
              audit=False)                                        # VAT charged on the TD for a week
    for i in range(4):                                            # TD posted twice on one night
        make_stay(m, f"T5:{i}", channel="WALKIN", check_in=AUG + timedelta(days=4 + 6 * i), nights=2, rate=_rate(m),
                  rate_plan="RACK", audit=False)
    for i in range(3):                                            # TD on a no-show
        make_stay(m, f"T6:{i}", channel="DIRECT_WEB", check_in=AUG + timedelta(days=6 + 7 * i), nights=1 + i % 2,
                  rate=_rate(m), status="NO_SHOW", rate_plan="BAR_DIRECT", audit=False)


def td_unit(m: Money) -> int:
    """The TD per unit-night in force for the plant nights (AED 10 for a
    Dubai 3-star), read from the night the PMS priced."""
    m.cur.execute("SELECT td_minor FROM reservation_nights WHERE reservation_id = %s ORDER BY stay_date LIMIT 1",
                  (m.stays["T3"]["id"],))
    return int(m.cur.fetchone()["td_minor"])


def folio_export_rows(m: Money, unit: int) -> list[dict]:
    """The PMS folio of the T stays as exported: the night audit's postings
    with the planted mistakes."""
    s = m.stays
    tkeys = ["T1", "T2:a", "T2:b", "T3", "T4"] + [f"T5:{i}" for i in range(4)]
    rows = folio_rows(m, [s[k]["id"] for k in tkeys])
    by_pms = {s[k]["pms_res_id"]: k for k in tkeys}
    out: list[dict] = []
    for r in rows:
        k = by_pms[r["pms_res_id"]]
        if r["code"] == "TD" and k == "T1":
            r["amount_minor"] = unit                      # one unit posted for three rooms
        if r["code"] == "TD" and k == "T3":
            r["amount_minor"] = T3_POSTED
        if r["code"] == "VAT" and k == "T4":
            r["amount_minor"] += pct_of(unit, Decimal(5))  # VAT on the TD
        out.append(r)
        if r["code"] == "TD" and k.startswith("T5:") and r["stay_date"] == s[k]["check_in"]:
            out.append(r | {"pms_txn_id": r["pms_txn_id"] + "-2", "description": "Tourism Dirham (reposted)"})
    b = s["T2:b"]
    posted = {r["stay_date"] for r in out if r["code"] == "TD" and r["pms_res_id"] == b["pms_res_id"]}
    for i in range(b["nights"]):                          # TD kept on past night 30 of the chain
        d = b["check_in"] + timedelta(days=i)
        if d not in posted:
            out.append({"pms_txn_id": f"{b['pms_res_id']}-{d:%Y%m%d}-TD", "posted_date": d, "stay_date": d,
                        "code": "TD", "amount_minor": unit, "pms_res_id": b["pms_res_id"], "room_number": None,
                        "description": "Td", "qty": 1})
    for i in range(3):                                    # the night audit charged TD on no-shows
        st = s[f"T6:{i}"]
        for n in range(st["nights"]):
            d = st["check_in"] + timedelta(days=n)
            out.append({"pms_txn_id": f"{st['pms_res_id']}-{d:%Y%m%d}-TD", "posted_date": d, "stay_date": d,
                        "code": "TD", "amount_minor": unit, "pms_res_id": st["pms_res_id"], "room_number": None,
                        "description": "Td", "qty": 1})
    return sorted(out, key=lambda r: (r["posted_date"], r["pms_txn_id"]))


FOLIO_HEADERS = ("Transaction ID", "Posting date", "Transaction code", "Amount", "Confirmation no", "Stay date",
                 "Description", "Quantity", "Currency", "Room")


def import_folio(m: Money) -> int:
    unit = td_unit(m)
    rows = folio_export_rows(m, unit)
    body: list[list[Any]] = [list(FOLIO_HEADERS)]
    for r in rows:
        body.append([r["pms_txn_id"], r["posted_date"].isoformat(), r["code"], dot(r["amount_minor"]), r["pms_res_id"],
                     r["stay_date"].isoformat(), r["description"], r["qty"], CCY, r["room_number"] or ""])
    body.append(["Total", "", "", dot(sum(r["amount_minor"] for r in rows)), "", "", "", "", "", ""])
    with at(m.ctx, m.ctx.anchor - timedelta(days=1), 6, 30):
        import_file(m, "folio", (_csv(body).encode("utf-8"), "pms-folio-export.csv", "text/csv"), kind="PMS_FOLIO",
                    source_code="pms_folio")
    return unit


def tax_ledger(m: Money) -> date:
    """The owner applied the announced TD exemption to one night (the PMS
    kept posting TD: TX1); each month closed on the 1st (the rollup job),
    then was remitted on the 10th: every month exactly except August, AED
    350 short (T7). Returns the exempt night."""
    a = m.ctx.anchor
    exempt = a - timedelta(days=EXEMPT_NIGHT_OFFSET)
    core_tax.set_override(m.cur, m.slug, TD_RULE, apply=True, effective_from=exempt, effective_to=exempt,
                          note="Owner applied the announced exemption for one night")
    month = JUN
    while True:
        close = month_end(month) + timedelta(days=1)
        if close > a:
            break
        with at(m.ctx, close, 4, 5):
            run_job(m, "tax.td_ledger_rollup")
        # Remitted on the 10th, before T7's 5-day window before the 15th;
        # a month whose 10th is still ahead stays open (due in > 5 days).
        closed, month = month, close
        paid_on, short = close + timedelta(days=9), TD_SHORT_MINOR if closed == AUG else 0
        if paid_on > a:
            continue
        m.cur.execute(
            """
            UPDATE pfg_td_ledger SET td_remitted_minor = td_collected_minor - %s, remitted_on = %s,
                   remit_reference = %s, status = CASE WHEN %s > 100 THEN 'MISMATCH' ELSE 'REMITTED' END,
                   updated_at = NOW()
             WHERE hotel_slug = %s AND month = %s
            """,
            (short, paid_on, f"DET-TD-{closed:%Y%m}", short, m.slug, closed),
        )
    return exempt


def tax_plants(m: Money, unit: int, exempt: date) -> None:
    s = m.stays

    def one(det: str, key: str, amount: int, month: date, note: str, klass: str = "COMPLIANCE") -> None:
        st = s[key]
        plant(m, det, f"{det}:{st['pms_res_id']}", klass=klass, expected_minor=amount, fingerprints=[],
              subjects=[{"pms_res_id": st["pms_res_id"], "month": mkey(month)}], ids=_ids(st), note=note)

    one("T1", "T1", 2 * unit * 4, AUG, "group of 3 rooms x 4 nights, TD posted once a night")
    one("T2", "T2:b", 15 * unit, SEP, "25 + 20 nights back to back: TD kept on nights 31-45")
    one("T3", "T3", 10 * abs(T3_POSTED - unit), AUG, "TD keyed at AED 15 for 10 nights")
    one("T4", "T4", pct_of(7 * unit, Decimal(5)), AUG, "a week of VAT charged on a walk-in's TD")
    for i in range(4):
        one("T5", f"T5:{i}", unit, AUG, "TD posted twice on one night", klass="RECOVERABLE")
    for i in range(3):
        one("T6", f"T6:{i}", unit * s[f"T6:{i}"]["nights"], AUG, "TD charged on a no-show")
    plant(m, "T7", f"T7:{mkey(AUG)}", klass="COMPLIANCE", expected_minor=TD_SHORT_MINOR,
          fingerprints=[fingerprint("T7", None, "ledger", mkey(AUG))], ids={"month": mkey(AUG)},
          note="August TD remitted AED 350 below what the folio collected")
    m.cur.execute(
        """
        SELECT f.reservation_id AS rid, r.pms_res_id,
               SUM(CASE WHEN f.is_reversal THEN -ABS(f.amount_minor) ELSE f.amount_minor END) AS td
          FROM folio_charges f JOIN reservations r ON r.id = f.reservation_id
         WHERE f.hotel_slug = %s AND f.code = 'TD' AND COALESCE(f.stay_date, f.posted_date) = %s
           AND r.status NOT IN ('CANCELLED', 'NO_SHOW')
           AND EXISTS (SELECT 1 FROM reservation_nights n WHERE n.reservation_id = r.id AND n.stay_date = %s
                                                            AND n.stayed)
         GROUP BY 1, 2 HAVING SUM(CASE WHEN f.is_reversal THEN -ABS(f.amount_minor) ELSE f.amount_minor END) > 0
         ORDER BY 2
        """,
        (m.slug, exempt, exempt),
    )
    for r in m.cur.fetchall():
        plant(m, "TX1", f"TX1:{r['pms_res_id']}", klass="COMPLIANCE", expected_minor=int(r["td"]), fingerprints=[],
              subjects=[{"pms_res_id": r["pms_res_id"], "month": mkey(exempt)}],
              ids={"pms_res_ids": [r["pms_res_id"]], "exempt_night": exempt.isoformat()},
              note="TD posted on a night inside the exemption the owner applied")


# ─── Generated files ───


def write_files(m: Money) -> str | None:
    """With CLEFS_DEMO_FILES_DIR set, write every generated document there
    (the canonical set lives in seed/demo_uae/files/, for trying the upload
    screens and the public audit by hand)."""
    target = os.environ.get(FILES_ENV)
    if not target:
        return None
    os.makedirs(target, exist_ok=True)
    for name, data in sorted(m.files.items()):
        with open(os.path.join(target, name), "wb") as fh:
            fh.write(data)
    return target


# ─── Part entry ───


def seed_part(ctx: Any) -> dict:
    """Plant the money cases and run detection. Skipped (with a reason)
    when Profit Guard is off for the hotel or the anchor is before the
    August statement could have arrived and been checked."""
    api = ctx.module_api("profit_guard")
    if api is None or not ctx.table_exists("pfg_ota_statements"):
        return {"skipped": "profit_guard is not enabled"}
    if ctx.anchor < EARLIEST_ANCHOR:
        return {"skipped": f"anchor before {EARLIEST_ANCHOR.isoformat()}"}
    from clefs import registry

    registry.load_module("profit_guard")  # commit handlers, post-commit hooks, jobs and tasks
    m = Money(ctx=ctx, api=api, rng=ctx.rng("money:plants"))
    m.cur.execute("SELECT ota_res_id FROM reservations WHERE hotel_slug = %s AND ota_res_id IS NOT NULL", (ctx.slug,))
    m.ota_ids = {r["ota_res_id"] for r in m.cur.fetchall()}
    scope_genius(m)
    for make in (statement_stays, payout_stays, card_stays, l14_stays, tax_stays):
        make(m)
    m.cur.execute("ANALYZE reservations, reservation_nights, folio_charges")
    import_statements(m)
    import_credit_note(m)
    import_payouts(m)
    import_cards(m)
    unit = import_folio(m)
    rate_evidence(m)
    exempt = tax_ledger(m)
    # The anchor evening: the full scan (every detector, auto-close) and the
    # nightly tax checks (last and this month, plus T7 over the ledger).
    detect(m)
    run_job(m, "tax.td_checks")
    l14_plants(m)
    l15_plants(m)
    tax_plants(m, unit, exempt)
    written = write_files(m)
    ctx.stats["money"] = {"files": sorted(m.files), "batches": {k: b["id"] for k, b in m.batches.items()}}
    out = dict(sorted(m.counts.items())) | {"stays": len(m.stays), "files": len(m.files),
                                             "plants": sum(1 for p in ctx.plants if p["klass"] != "DECOY")}
    return out | ({"files_dir": written} if written else {})


# ─── demo-check ───


def money_checks(cur: Any, slug: str) -> list[tuple[str, bool, str]]:
    """Against the manifest the CLI wrote (expected_plants.json): every
    planted finding exists, no decoy fired, claims are >= 90% plants, and
    RECOVERABLE clears 3 x an AED 1,500 subscription. Read-only."""
    from clefs.seed.demo_uae import MANIFEST_PATH

    try:
        with open(MANIFEST_PATH, encoding="utf-8") as fh:
            manifest = json.load(fh)
    except (OSError, ValueError):
        return [("money: manifest", False, f"{MANIFEST_PATH} not readable")]
    if manifest.get("hotel") != slug:
        return []  # the manifest describes another demo hotel (a test or --manifest elsewhere)
    cur.execute("SELECT to_regclass('public.pfg_findings') IS NOT NULL AS ok")
    if not cur.fetchone()["ok"]:
        return [("money: findings", False, "pfg_findings missing")]
    cur.execute("SELECT fingerprint, class, amount_minor, currency, internal_only FROM pfg_findings "
                "WHERE hotel_slug = %s AND source = 'detector'", (slug,))
    rows = {r["fingerprint"]: dict(r) for r in cur.fetchall()}
    plants = [p for p in manifest.get("plants", []) if p.get("klass") != "DECOY"]
    want = {fp for p in plants for fp in fingerprints_of(cur, slug, p)}
    decoys = {fp for p in manifest.get("plants", []) for fp in p.get("decoy_fingerprints", [])}
    missing = sorted(want - set(rows))
    fired = sorted(decoys & set(rows))
    claims = [fp for fp, r in rows.items() if r["class"] != "STRUCTURAL"]
    precision = sum(fp in want for fp in claims) / len(claims) if claims else 0.0
    recoverable = sum(int(r["amount_minor"]) for r in rows.values()
                      if r["class"] == "RECOVERABLE" and r["currency"] == CCY and not r["internal_only"])
    return [
        ("money: planted findings present", not missing, f"{len(want) - len(missing)}/{len(want)}"),
        ("money: decoys silent", not fired, f"{len(fired)} decoy finding(s)"),
        ("money: precision >= 0.9", precision >= 0.9, f"{precision:.3f} over {len(claims)} claims"),
        ("money: recoverable >= 3 x AED 1,500", recoverable >= 450000, f"AED {major(recoverable)}"),
    ]


CHECKS = [money_checks]
